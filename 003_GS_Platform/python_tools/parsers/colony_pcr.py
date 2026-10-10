"""Parser for Colony PCR Excel files.

Reads sheets matching "Plate {N}". For each plate:
- Grid data from B3:M10 (8 rows x 12 cols)
- Sample metadata from the table below the grid (JobID, BBID, Vector, etc.)

Returns { success, message, data: [{ plateId, sheetName, gridData, gridHeaders, samplesData, samplesHeaders }] }.
"""

import openpyxl
import re
from .base import ROW_LABELS


def parse(file_bytes: bytes) -> dict:
    try:
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)

        plate_sheet_names = [name for name in wb.sheetnames if re.match(r'^plate \d+$', name, re.IGNORECASE)]

        if len(plate_sheet_names) == 0:
            return {"success": False, "message": 'No sheets named "Plate {number}" found in the file.'}

        plates = []
        grid_headers = ['Row'] + [f'C{i}' for i in range(1, 13)]
        notes = []

        for sheet_name in plate_sheet_names:
            ws = wb[sheet_name]
            match = re.search(r'\d+', sheet_name)
            plate_id = int(match.group()) if match else 0

            # Read grid data from B3:M10 (row 3-10, col 2-13)
            grid_json = []
            for row_idx, row_label in enumerate(ROW_LABELS):
                row_obj = {"Row": row_label}
                for col_idx in range(12):
                    cell_val = ws.cell(row=3 + row_idx, column=2 + col_idx).value
                    row_obj[f'C{col_idx + 1}'] = str(cell_val) if cell_val is not None else ''
                grid_json.append(row_obj)

            # Find samples table: look for row where col B='JobID' and col C='BBID'.
            # Hardened (v0.0.5 Phase 1): if the header scan or the table read
            # hits anything unexpected, keep the plate grid and leave the
            # samples table empty instead of raising — one bad sheet must not
            # fail the whole file (P1_IMPORT_SPEC.md §3.2).
            sample_header_row = None
            samples_data = []
            samples_headers = []
            try:
                for row_idx in range(1, ws.max_row + 1):
                    b_val = ws.cell(row=row_idx, column=2).value
                    c_val = ws.cell(row=row_idx, column=3).value
                    if b_val == 'JobID' and c_val == 'BBID':
                        sample_header_row = row_idx
                        break

                if sample_header_row is not None:
                    # Read headers from columns B-G (cols 2-7)
                    raw_headers = []
                    for col in range(2, 8):
                        h = ws.cell(row=sample_header_row, column=col).value
                        if h is not None:
                            raw_headers.append(str(h))
                    samples_headers = raw_headers

                    # Read data rows until empty
                    for data_row in range(sample_header_row + 1, ws.max_row + 1):
                        first_val = ws.cell(row=data_row, column=2).value
                        if not first_val:
                            break
                        sample_row = {}
                        for h_idx, header in enumerate(samples_headers):
                            cell_val = ws.cell(row=data_row, column=2 + h_idx).value
                            sample_row[header] = str(cell_val) if cell_val is not None else ''
                        samples_data.append(sample_row)
            except Exception as exc:
                samples_data = []
                samples_headers = []
                notes.append(
                    f"{sheet_name}: samples table read failed ({exc}); grid kept, samples left empty."
                )

            if sample_header_row is None:
                notes.append(
                    f"{sheet_name}: no JobID/BBID header row found — grid kept, samples left empty."
                )

            plates.append({
                "plateId": plate_id,
                "sheetName": sheet_name,
                "gridData": grid_json,
                "gridHeaders": grid_headers,
                "samplesData": samples_data,
                "samplesHeaders": samples_headers,
            })

        return {"success": True, "message": f"Found {len(plates)} plate(s).", "data": plates, "notes": notes}

    except Exception as e:
        return {"success": False, "message": str(e) or "Error parsing Colony PCR file."}
