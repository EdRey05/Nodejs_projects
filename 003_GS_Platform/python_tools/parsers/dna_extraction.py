"""Parser for DNA Extraction Excel files.

Reads sheet "Layout":
- Plate 1: range B5:G12 (8 rows x 6 cols)
- Plate 2: range J5:O12 (8 rows x 6 cols)

Returns { success, message, data: [{ plateId, data: { wellId: { id, sample, params, content } } }] }.
"""

import openpyxl
from .base import ROW_LABELS, generate_plate_data


def parse(file_bytes: bytes) -> dict:
    try:
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)

        sheet_name = 'Layout'
        if sheet_name not in wb.sheetnames:
            return {"success": False, "message": f'Sheet "{sheet_name}" not found.'}

        ws = wb[sheet_name]

        # Define the two plate ranges
        ranges = [
            {"id": 1, "min_row": 5, "min_col": 2, "max_col": 7},   # B5:G12
            {"id": 2, "min_row": 5, "min_col": 10, "max_col": 15},  # J5:O12
        ]

        plates = []
        for r in ranges:
            # Read 8 rows x 6 cols
            data = []
            for row_idx in range(8):
                row_vals = []
                for col_idx in range(6):
                    cell_val = ws.cell(
                        row=r["min_row"] + row_idx,
                        column=r["min_col"] + col_idx
                    ).value
                    row_vals.append(cell_val)
                data.append(row_vals)

            # Check if range is entirely empty
            is_empty = all(
                cell is None or str(cell).strip() == ''
                for row in data for cell in row
            )
            if is_empty:
                continue

            plate_data = generate_plate_data(8, 6)
            for i in range(8):
                row_char = ROW_LABELS[i]
                for j in range(6):
                    well_id = f"{row_char}{j + 1}"
                    cell_val = data[i][j]
                    plate_data[well_id]["content"] = str(cell_val).strip() if cell_val is not None else ''

            plates.append({
                "plateId": r["id"],
                "data": plate_data,
            })

        if len(plates) == 0:
            return {
                "success": False,
                "message": 'No data found in the expected ranges (B5:G12, J5:O12) of the "Layout" sheet.',
            }

        return {"success": True, "message": f"Found data for {len(plates)} plate(s).", "data": plates}

    except Exception as e:
        return {"success": False, "message": str(e) or "Error parsing DNA Extraction file."}
