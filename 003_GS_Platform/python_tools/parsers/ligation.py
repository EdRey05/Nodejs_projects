"""Parser for Ligation Excel files.

Reads sheet "Plate Name", range A1:W97. Filters rows that have Name or BBID.
Returns { success, message, data: [{header: value, ...}], headers: [str] }.
"""

import openpyxl


def parse(file_bytes: bytes) -> dict:
    try:
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)

        sheet_name = 'Plate Name'
        if sheet_name not in wb.sheetnames:
            return {"success": False, "message": f'Sheet "{sheet_name}" not found.'}

        ws = wb[sheet_name]

        # Read range A1:W97 (rows 1-97, cols 1-23)
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=97, min_col=1, max_col=23, values_only=True):
            rows.append(list(row))

        if len(rows) < 2:
            return {"success": False, "message": "No data rows found."}

        original_headers = rows[0]
        headers = [str(h) for h in original_headers if h is not None and str(h).strip() != '']

        data_rows = rows[1:]

        # Find Name and BBID column indices
        name_index = None
        bbid_index = None
        for i, h in enumerate(original_headers):
            if h == 'Name':
                name_index = i
            elif h == 'BBID':
                bbid_index = i

        if name_index is None or bbid_index is None:
            return {"success": False, "message": 'Missing "Name" and/or "BBID" columns.'}

        # Filter rows that have Name or BBID
        filtered_rows = [r for r in data_rows if r[name_index] or r[bbid_index]]
        if len(filtered_rows) == 0:
            return {"success": False, "message": "No valid rows found to import."}

        data_to_return = []
        for row in filtered_rows:
            row_data = {}
            for header in headers:
                header_index = next(
                    (i for i, h in enumerate(original_headers) if h == header), None
                )
                if header_index is not None:
                    val = row[header_index]
                    row_data[header] = val if val is not None else None
            data_to_return.append(row_data)

        return {
            "success": True,
            "message": f"{len(filtered_rows)} rows ready.",
            "data": data_to_return,
            "headers": headers,
        }

    except Exception as e:
        return {"success": False, "message": str(e) or "Error parsing ligation file."}
