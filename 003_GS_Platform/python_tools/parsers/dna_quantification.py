"""Parser for DNA Quantification Excel files.

Reads the first sheet. Extracts headers from row 0 (cols 1-6), then all data rows.
Returns { success, message, data: [{header: value, ...}], headers: [str] }.
"""

import openpyxl


def parse(file_bytes: bytes) -> dict:
    try:
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)

        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]

        # Read all rows as lists
        all_rows = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(list(row))

        if len(all_rows) < 2:
            return {"success": False, "message": "No data found."}

        # Headers from first row, columns index 1-6 (B-G)
        all_headers = all_rows[0]
        headers = []
        for i in range(1, 7):
            if i < len(all_headers) and all_headers[i] is not None:
                headers.append(str(all_headers[i]))

        data_to_return = []
        for row in all_rows[1:]:
            row_data = {}
            for idx, header in enumerate(headers):
                col_idx = idx + 1
                val = row[col_idx] if col_idx < len(row) else None
                row_data[header] = val if val is not None else ''
            # Filter out completely empty rows
            if any(v != '' and v is not None for v in row_data.values()):
                data_to_return.append(row_data)

        return {
            "success": True,
            "message": f"{len(data_to_return)} rows parsed successfully.",
            "data": data_to_return,
            "headers": headers,
        }

    except Exception as e:
        return {"success": False, "message": str(e) or "Error parsing DNA Quantification file."}
